from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


def export_excel(classified: pd.DataFrame, insights: dict[str, pd.DataFrame], path: str | Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        insights["kpis"].to_excel(writer, sheet_name="KPI Summary", index=False)
        insights["topics"].to_excel(writer, sheet_name="Topic Trends", index=False)
        insights["sentiment"].to_excel(writer, sheet_name="Sentiment", index=False)
        insights["priority"].to_excel(writer, sheet_name="Priority", index=False)
        insights["opportunities"].to_excel(writer, sheet_name="Opportunities", index=False)
        classified.to_excel(writer, sheet_name="Classified Reviews", index=False)
        for sheet in writer.sheets.values():
            sheet.freeze_panes = "A2"
            for column in sheet.columns:
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 38)


def export_powerpoint(insights: dict[str, pd.DataFrame], path: str | Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    _title_slide(prs, "Customer Feedback Intelligence", "AI-assisted review analysis | Executive readout")
    kpis = insights["kpis"].set_index("KPI")["Value"].to_dict()
    _bullets(prs, "Executive snapshot", [
        f"{kpis['Total reviews']} reviews analyzed with an average rating of {kpis['Average rating']}/5.",
        f"Positive rate: {kpis['Positive rate']:.0%}; negative rate: {kpis['Negative rate']:.0%}.",
        f"{kpis['High-priority rate']:.0%} of reviews require urgent product or service attention.",
    ])
    top = insights["topics"].head(3)
    _bullets(prs, "What customers talk about", [
        f"{row.topic}: {row.reviews} reviews ({row.share:.0%} share), average rating {row.avg_rating:.1f}/5."
        for row in top.itertuples()
    ])
    actions = insights["opportunities"].head(4)
    _bullets(prs, "Recommended actions", [
        f"{row.topic}: {row.issue_or_opportunity} ({row.reviews} related reviews)."
        for row in actions.itertuples()
    ])
    _bullets(prs, "Sales and service implications", [
        "Equip sales teams with proof points around quality, speed, and onboarding.",
        "Use high-priority issues as service-recovery triggers before renewal or expansion conversations.",
        "Close the loop: publish owners, target dates, and a monthly feedback KPI review.",
    ])
    prs.save(path)


def _title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    for shape in (slide.shapes.title, slide.placeholders[1]):
        shape.text_frame.word_wrap = True
        shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    slide.placeholders[1].text_frame.paragraphs[0].font.size = Pt(18)


def _bullets(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.word_wrap = True
    slide.shapes.title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    box = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(11.85), Inches(5.35))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.15)
    frame.margin_right = Inches(0.15)
    frame.margin_top = Inches(0.1)
    frame.margin_bottom = Inches(0.1)
    for index, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        paragraph.space_after = Pt(12)
